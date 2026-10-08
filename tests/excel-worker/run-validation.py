"""Phase 2A validation: constrained, copy-on-write workbook transformation."""
from __future__ import annotations
import base64
import json
import os
import tempfile
from pathlib import Path
import sys
sys.path.insert(0, "/app")
from openpyxl import Workbook, load_workbook

os.environ["EXCEL_STORAGE_ROOT"] = tempfile.mkdtemp(prefix="excel-phase2a-")
os.environ["EXCEL_INTERNAL_API_TOKEN"] = "test-token"
os.environ["EXCEL_MAX_ROWS"] = "100"
os.environ["EXCEL_MAX_COLUMNS"] = "20"
import app

root = Path(os.environ["EXCEL_STORAGE_ROOT"])
imports = root / "imports"
imports.mkdir(parents=True, exist_ok=True)
source = imports / "source.xlsx"
book = Workbook()
sheet = book.active
sheet.title = "Bron"
sheet.append(["Naam", "Aantal"])
sheet.append(["Alpha", 1])
book.save(source)
book.close()

metadata = app.transform_workbook(source, {
    "filename": "getransformeerd.xlsx",
    "operations": [
        {"type": "rename_sheet", "sheet": "Bron", "newName": "Overzicht"},
        {"type": "append_rows", "sheet": "Overzicht", "rows": [["Beta", 2], ["Gamma", 3]]},
        {"type": "set_cell", "sheet": "Overzicht", "cell": "C1", "value": "Status"},
        {"type": "set_cell", "sheet": "Overzicht", "cell": "C2", "value": "Actief"},
        {"type": "format_range", "sheet": "Overzicht", "range": "A1:C1", "style": "limburg_header"},
        {"type": "freeze_panes", "sheet": "Overzicht", "cell": "A2"},
        {"type": "set_auto_filter", "sheet": "Overzicht", "range": "A1:C4"},
        {"type": "add_sheet", "sheet": "Notities"},
        {"type": "set_cell", "sheet": "Notities", "cell": "A1", "value": "Wijzigingen zijn veilig gekopieerd."},
    ],
}, "507f1f77bcf86cd799439011", "conversation-1")

target = app.FILES_DIR / metadata["storageName"]
assert source.read_bytes() != target.read_bytes(), "copy-on-write must create a new workbook"
book = load_workbook(target, data_only=False)
try:
    sheet = book["Overzicht"]
    assert sheet.max_row == 4 and sheet["A4"].value == "Gamma"
    assert sheet["C2"].value == "Actief"
    assert sheet["A1"].fill.fgColor.rgb.endswith("002F6C")
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:C4"
    assert book["Notities"]["A1"].value.startswith("Wijzigingen")
finally:
    book.close()
assert metadata["sourceSha256"] == app.hash_file(source)
assert metadata["ownerUserId"] == "507f1f77bcf86cd799439011"

for invalid in [
    {"operations": [{"type": "set_cell", "sheet": "Bron", "cell": "A0", "value": "x"}]},
    {"operations": [{"type": "set_cell", "sheet": "Ontbreekt", "cell": "A1", "value": "x"}]},
    {"operations": [{"type": "append_rows", "sheet": "Bron", "rows": "geen-lijst"}]},
    {"operations": [{"type": "delete_sheet", "sheet": "Bron"}]},
    {"operations": [{"type": "arbitrary_python", "code": "import os"}]},
]:
    try:
        app.transform_workbook(source, invalid, "user-1", "conversation-1")
    except app.ClientError:
        continue
    raise AssertionError(f"unsafe request accepted: {invalid}")

# The public transform endpoint accepts only base64 workbook bytes; no caller path is trusted.
encoded = base64.b64encode(source.read_bytes()).decode("ascii")
response = app.transform_excel({
    "userId": "507f1f77bcf86cd799439011",
    "conversationId": "conversation-1",
    "sourceFileName": "source.xlsx",
    "sourceFileBase64": encoded,
    "filename": "endpoint-copy.xlsx",
    "operations": [{"type": "set_cell", "sheet": "Bron", "cell": "B2", "value": 99}],
})
assert response["ok"] and response["fileName"] == "endpoint-copy.xlsx"
try:
    app.transform_excel({"userId": "user", "conversationId": "c", "sourcePath": "/etc/passwd", "operations": [{"type": "add_sheet", "sheet": "X"}]})
except Exception:
    pass
else:
    raise AssertionError("path-based transform request was accepted")

print(json.dumps({"status": "phase2a transform validation passed", "artifactId": metadata["artifactId"]}))
