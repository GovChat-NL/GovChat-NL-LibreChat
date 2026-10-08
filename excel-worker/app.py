"""Private Excel MVP service for GovChat.

Only the n8n workflow may call this service. User and conversation identity are
trusted only when injected by n8n from LibreChat bridge headers.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import secrets
import shutil
import tempfile
import uuid
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

APP_TOKEN = os.environ.get("EXCEL_INTERNAL_API_TOKEN", "")
STORAGE_ROOT = Path(os.environ.get("EXCEL_STORAGE_ROOT", "/storage"))
FILES_DIR = STORAGE_ROOT / "files"
METADATA_DIR = STORAGE_ROOT / "metadata"
MAX_UPLOAD_BYTES = int(os.environ.get("EXCEL_MAX_UPLOAD_BYTES", "10485760"))
MAX_SHEETS = int(os.environ.get("EXCEL_MAX_SHEETS", "25"))
MAX_ROWS = int(os.environ.get("EXCEL_MAX_ROWS", "10000"))
MAX_COLUMNS = int(os.environ.get("EXCEL_MAX_COLUMNS", "100"))
ARTIFACT_TTL_HOURS = int(os.environ.get("EXCEL_ARTIFACT_TTL_HOURS", "24"))
HEADER_FILL = "002F6C"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

app = FastAPI(title="GovChat Excel Worker", docs_url=None, redoc_url=None)


class ClientError(Exception):
    """A validation failure safe to show in the chat UI."""


def require_service_token(authorization: str | None = Header(default=None)) -> None:
    expected = f"Bearer {APP_TOKEN}"
    if not APP_TOKEN or not authorization or not secrets.compare_digest(authorization, expected):
        raise HTTPException(status_code=401, detail="Niet geautoriseerd voor Excel-verwerking.")


def fail(message: str, status_code: int = 400) -> HTTPException:
    return HTTPException(status_code=status_code, detail=message)


def normalize_text(value: Any, maximum: int = 200) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[\x00-\x1f\x7f]", " ", text).strip()
    return text[:maximum]


def safe_filename(value: str, fallback: str = "provincie-limburg-tabel.xlsx") -> str:
    filename = Path(normalize_text(value, 120)).name
    filename = re.sub(r"[^A-Za-z0-9À-ÿ._ -]+", "-", filename).strip(" .-")
    if not filename:
        filename = fallback
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    return filename[:120]


def valid_identity(user_id: str, conversation_id: str) -> tuple[str, str]:
    user = normalize_text(user_id, 128)
    conversation = normalize_text(conversation_id, 128)
    if not user or not conversation:
        raise ClientError("De geverifieerde gebruiker- of gesprekssessie ontbreekt.")
    return user, conversation


def check_xlsx_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 5000:
                raise ClientError("Het Excel-bestand bevat te veel onderdelen.")
            uncompressed = sum(item.file_size for item in infos)
            compressed = sum(item.compress_size for item in infos)
            if uncompressed > MAX_UPLOAD_BYTES * 20 or (compressed and uncompressed / compressed > 100):
                raise ClientError("Het Excel-bestand is onveilig groot om te inspecteren.")
    except zipfile.BadZipFile as exc:
        raise ClientError("Het bestand is geen geldig, niet-versleuteld XLSX-bestand.") from exc


def ensure_supported_upload(upload: UploadFile, path: Path, byte_count: int) -> None:
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ClientError("Alleen .xlsx- en .xlsm-bestanden worden in deze MVP ondersteund.")
    if byte_count == 0:
        raise ClientError("Het geüploade Excel-bestand is leeg.")
    if byte_count > MAX_UPLOAD_BYTES:
        raise ClientError("Het Excel-bestand is groter dan de toegestane limiet.")
    check_xlsx_zip(path)


def safe_cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return normalize_text(value, 300) if isinstance(value, str) else value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return normalize_text(value, 300)


def inspect_workbook(path: Path) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:  # openpyxl error messages may disclose content; do not expose them.
        raise ClientError("Het Excel-bestand is beschadigd, versleuteld of niet leesbaar.") from exc

    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise ClientError("Het Excel-bestand bevat meer werkbladen dan toegestaan.")
        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
                raise ClientError("Een werkblad overschrijdt de toegestane rij- of kolomlimiet.")
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, ())
            columns = [normalize_text(value, 120) or f"Kolom {index + 1}" for index, value in enumerate(header_row)]
            previews: list[dict[str, Any]] = []
            for row in rows:
                if len(previews) == 3:
                    break
                preview = {
                    columns[index] if index < len(columns) else f"Kolom {index + 1}": safe_cell_value(value)
                    for index, value in enumerate(row[:MAX_COLUMNS])
                }
                if any(value is not None for value in preview.values()):
                    previews.append(preview)
            sheets.append({"name": normalize_text(worksheet.title, 80), "columns": columns, "sampleRows": previews})
        return {"sheets": sheets}
    finally:
        workbook.close()


def validate_table(payload: dict[str, Any]) -> tuple[list[str], list[list[Any]], str, str]:
    headers_raw = payload.get("headers")
    rows_raw = payload.get("rows")
    if not isinstance(headers_raw, list) or not headers_raw or len(headers_raw) > MAX_COLUMNS:
        raise ClientError("Geef 1 tot en met 100 kolomkoppen op.")
    if not isinstance(rows_raw, list) or len(rows_raw) > MAX_ROWS:
        raise ClientError("Geef maximaal 10.000 tabelrijen op.")
    headers = [normalize_text(value, 120) or f"Kolom {index + 1}" for index, value in enumerate(headers_raw)]
    rows: list[list[Any]] = []
    for item in rows_raw:
        if isinstance(item, dict):
            row = [safe_cell_value(item.get(header)) for header in headers]
        elif isinstance(item, list):
            row = [safe_cell_value(value) for value in item[: len(headers)]]
            row.extend([None] * (len(headers) - len(row)))
        else:
            raise ClientError("Elke tabelrij moet een lijst of object zijn.")
        rows.append(row)
    sheet_name = normalize_text(payload.get("sheetName", "Tabel"), 31) or "Tabel"
    filename = safe_filename(str(payload.get("filename", "provincie-limburg-tabel.xlsx")))
    return headers, rows, sheet_name, filename


def create_artifact(payload: dict[str, Any], user_id: str, conversation_id: str) -> dict[str, Any]:
    headers, rows, sheet_name, display_name = validate_table(payload)
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    METADATA_DIR.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    fill = PatternFill("solid", fgColor=HEADER_FILL)
    font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="B7C9E2"), right=Side(style="thin", color="B7C9E2"),
        top=Side(style="thin", color="B7C9E2"), bottom=Side(style="thin", color="B7C9E2"),
    )
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.border = border
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, header in enumerate(headers, start=1):
        longest = max([len(str(header)), *(len(normalize_text(row[index - 1], 60)) for row in rows)] or [10])
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 40)

    artifact_id = str(uuid.uuid4())
    storage_name = f"{artifact_id}.xlsx"
    target = FILES_DIR / storage_name
    with tempfile.NamedTemporaryFile(dir=FILES_DIR, suffix=".xlsx", delete=False) as temporary:
        temp_path = Path(temporary.name)
    try:
        workbook.save(temp_path)
        digest = hashlib.sha256(temp_path.read_bytes()).hexdigest()
        os.replace(temp_path, target)
    finally:
        workbook.close()
        temp_path.unlink(missing_ok=True)

    created_at = datetime.now(UTC)
    metadata = {
        "artifactId": artifact_id,
        "storageName": storage_name,
        "displayName": display_name,
        "mimeType": XLSX_MIME,
        "ownerUserId": user_id,
        "conversationId": conversation_id,
        "rowCount": len(rows),
        "sha256": digest,
        "createdAt": created_at.isoformat(),
        "expiresAt": (created_at + timedelta(hours=ARTIFACT_TTL_HOURS)).isoformat(),
    }
    metadata_path = METADATA_DIR / f"{artifact_id}.json"
    temporary_metadata = metadata_path.with_suffix(".json.tmp")
    temporary_metadata.write_text(json.dumps(metadata, separators=(",", ":")), encoding="utf-8")
    os.replace(temporary_metadata, metadata_path)
    return metadata


@app.get("/healthz")
def healthcheck() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/v1/inspect", dependencies=[Depends(require_service_token)])
async def inspect_excel(
    file: UploadFile = File(...),
    user_id: str = Form(...),
    conversation_id: str = Form(...),
) -> dict[str, Any]:
    try:
        valid_identity(user_id, conversation_id)
        with tempfile.NamedTemporaryFile(suffix=Path(file.filename or "").suffix, delete=False) as temporary:
            path = Path(temporary.name)
            copied = 0
            while chunk := await file.read(1024 * 1024):
                copied += len(chunk)
                if copied > MAX_UPLOAD_BYTES:
                    raise ClientError("Het Excel-bestand is groter dan de toegestane limiet.")
                temporary.write(chunk)
        try:
            ensure_supported_upload(file, path, copied)
            return {"ok": True, "fileName": safe_filename(file.filename or "bestand.xlsx"), **inspect_workbook(path)}
        finally:
            path.unlink(missing_ok=True)
    except ClientError as exc:
        raise fail(str(exc)) from exc


@app.post("/v1/generate", dependencies=[Depends(require_service_token)])
def generate_styled_excel(request: dict[str, Any]) -> dict[str, Any]:
    try:
        user_id, conversation_id = valid_identity(str(request.get("userId", "")), str(request.get("conversationId", "")))
        metadata = create_artifact(request, user_id, conversation_id)
        return {
            "ok": True,
            "artifactId": metadata["artifactId"],
            "fileName": metadata["displayName"],
            "rowCount": metadata["rowCount"],
            "createdAt": metadata["createdAt"],
            "downloadPath": f"/api/files/download/{metadata['artifactId']}",
        }
    except ClientError as exc:
        raise fail(str(exc)) from exc
